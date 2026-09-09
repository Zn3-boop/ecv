"""内容生成API - 调用LLM生成代码/文章/Office文档"""
from __future__ import annotations

import asyncio
import os
import traceback
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel, Field

from app.services.llm.provider import LLMProviderError, provider
from app.services.generator_config import config

router = APIRouter(prefix="/content", tags=["content_generator"])

CONTENT_LLM_TIMEOUT = float(os.getenv("CONTENT_LLM_TIMEOUT_SECONDS", "90"))


class GenerateRequest(BaseModel):
    """内容生成请求"""
    task: str = Field(..., description="生成任务描述，如'登录页面'、'毕业论文'")
    content_type: str = Field(default="html", description="内容类型：html/python/word/ppt/article/code/document")
    theme: str = Field(default="通用", description="主题或领域")
    length: str = Field(default="中等", description="长度：短/中/长")
    style: str = Field(default="现代", description="风格")
    output_path: str | None = Field(default=None, description="输出文件路径或目录")
    browser_url: str | None = Field(default=None, description="浏览器URL（用于打开生成的内容）")
    auto_open: bool = Field(default=False, description="生成后自动打开文件")


class GenerateResponse(BaseModel):
    """生成响应"""
    success: bool
    content: str = Field(default="", description="生成的内容")
    content_type: str
    output_file: str | None = None
    browser_opened: bool = False
    message: str
    error_detail: str | None = None  # 新增：暴露真实错误


async def call_llm(prompt: str) -> str:
    """调用已配置的LLM生成内容（带独立超时保护）"""
    if not provider.enabled:
        return "LLM调用失败: 未配置有效的AI Provider，请检查providers.json配置"
    
    try:
        response = await asyncio.wait_for(
            provider.chat(
                system_prompt="你是一个专业的内容生成助手，请根据用户要求生成高质量的内容。直接返回内容，不要解释。",
                user_prompt=prompt,
                temperature=0.7,
            ),
            timeout=CONTENT_LLM_TIMEOUT,
        )
        return response if response else "LLM调用失败: 返回内容为空"
    except asyncio.TimeoutError:
        return f"LLM调用失败: 生成超时（{CONTENT_LLM_TIMEOUT}s），请缩短内容长度或稍后重试"
    except LLMProviderError as e:
        return f"LLM调用失败: {str(e)}"
    except Exception as e:
        print(f"[LLM ERROR] {traceback.format_exc()}")
        return f"LLM调用异常: {str(e)}"


def save_file(path: str, content: str) -> tuple[str | None, str | None]:
    """保存内容到文件，返回 (文件路径, 错误信息)"""
    try:
        dir_path = os.path.dirname(path)
        if dir_path:
            os.makedirs(dir_path, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path, None
    except Exception as e:
        print(f"[SAVE ERROR] path={path}, err={traceback.format_exc()}")
        return None, str(e)


def _build_simple_rtf(content: str, title: str = "") -> str:
    def _esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}").replace("\n", "\\par\n")
    body = _esc(content.strip())
    t = _esc(title or "文档")
    return (
        r"{\rtf1\ansi\ansicpg936\deff0\nouicompat"
        r"{\fonttbl{\f0\fnil\fcharset134 SimSun;}}"
        r"\viewkind4\uc1\pard\lang2052\f0\fs24 "
        + (f"\\b {t}\\b0\\par\\par " if t else "")
        + body
        + r"\par}"
    )


def save_word_document(path: str, content: str, title: str = "") -> tuple[str | None, str | None]:
    """保存内容为Word文档(.docx)，失败则回退到.rtf（Word可打开，记事本不会抢开）"""
    try:
        dir_path = os.path.dirname(path)
        if dir_path:
            os.makedirs(dir_path, exist_ok=True)
        
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            
            doc = Document()
            
            _CN_FONT = '微软雅黑'
            style = doc.styles['Normal']
            style.font.name = _CN_FONT
            style.element.rPr.rFonts.set(qn('w:eastAsia'), _CN_FONT)
            
            lines = content.strip().split('\n')
            current_para = []
            
            def _set_cn_font(run, size=None, bold=False):
                run.font.name = _CN_FONT
                run.element.rPr.rFonts.set(qn('w:eastAsia'), _CN_FONT)
                if size:
                    run.font.size = Pt(size)
                if bold:
                    run.bold = True
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                if line.startswith('【') and line.endswith('】'):
                    if current_para:
                        p = doc.add_paragraph('\n'.join(current_para))
                        for r in p.runs:
                            _set_cn_font(r)
                        current_para = []
                    p = doc.add_paragraph(line[1:-1])
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        _set_cn_font(run, size=18, bold=True)
                elif any(line.startswith(prefix) for prefix in ['章节', '第', '一、', '二、', '三、', '四、', '五、']) or line in ['结论', '摘要', '引言']:
                    if current_para:
                        p = doc.add_paragraph('\n'.join(current_para))
                        for r in p.runs:
                            _set_cn_font(r)
                        current_para = []
                    p = doc.add_paragraph(line)
                    for run in p.runs:
                        _set_cn_font(run, size=14, bold=True)
                else:
                    current_para.append(line)
            
            if current_para:
                p = doc.add_paragraph('\n'.join(current_para))
                for r in p.runs:
                    _set_cn_font(r)
            
            doc.save(path)
            return path, None
        except ImportError:
            print("[WORD WARN] python-docx未安装，回退为RTF（Word可打开，记事本不会抢开）")
            rtf_path = path.replace('.docx', '.rtf')
            rtf_content = _build_simple_rtf(content, title)
            with open(rtf_path, "w", encoding="utf-8") as f:
                f.write(rtf_content)
            return rtf_path, "python-docx未安装，已保存为RTF格式（可用Word/WPS打开）"
    except Exception as e:
        print(f"[WORD ERROR] path={path}, content_len={len(content)}, err={traceback.format_exc()}")
        try:
            rtf_path = path.replace('.docx', '.rtf') if path.endswith('.docx') else path + '.rtf'
            rtf_content = _build_simple_rtf(content, title)
            with open(rtf_path, "w", encoding="utf-8") as f:
                f.write(rtf_content)
            return rtf_path, f"Word保存异常已回退RTF: {str(e)}"
        except Exception as e2:
            return None, f"保存完全失败: {str(e2)}"


def open_file(path: str) -> bool:
    """用系统默认程序打开文件"""
    try:
        import platform
        if platform.system() == 'Windows':
            os.startfile(path)
            return True
        else:
            import subprocess
            subprocess.Popen(['open', path])
            return True
    except Exception as e:
        print(f"[OPEN ERROR] {traceback.format_exc()}")
        return False


def open_in_browser(url: str) -> bool:
    """在浏览器中打开URL"""
    try:
        import subprocess
        subprocess.Popen(f'start "" "{url}"', shell=True)
        return True
    except Exception as e:
        print(f"[BROWSER ERROR] {traceback.format_exc()}")
        return False


def _resolve_output_path(req: GenerateRequest) -> tuple[str, str | None]:
    """
    解析输出路径，返回 (output_dir, output_file_path)
    如果 output_path 有扩展名，视为文件路径；否则视为目录
    """
    if req.output_path:
        if os.path.splitext(req.output_path)[1]:
            norm_path = os.path.normpath(req.output_path)
            return os.path.dirname(norm_path), norm_path
        else:
            return req.output_path.rstrip('/\\'), None
    else:
        default_dir = os.path.expanduser("~/Documents/AI_Generated")
        return default_dir, None


def _get_default_filename(content_type: str, file_path: str | None) -> str:
    """根据 content_type 返回默认文件名（带时间戳避免冲突）"""
    if file_path:
        return os.path.basename(file_path)
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    mapping = {
        "html": f"page_{ts}.html",
        "python": f"script_{ts}.py",
        "word": f"doc_{ts}.docx",
        "document": f"doc_{ts}.docx",
        "ppt": f"slides_{ts}.txt",
        "code": f"code_{ts}.py",
        "article": f"article_{ts}.txt",
    }
    return mapping.get(content_type, f"output_{ts}.txt")


@router.post("/generate", response_model=GenerateResponse)
async def generate_content(req: GenerateRequest):
    """生成各种类型的内容"""
    
    length_map = {"短": "500字", "中": "1000字", "长": "2000字以上"}
    length_desc = length_map.get(req.length, "1000字")
    
    # 构建 Prompt
    prompts = {
        "html": f"""请生成一个{req.style}风格的{req.task}HTML页面代码。
要求：
1. 完整的可保存HTML文件（包含<!DOCTYPE html>和<html>标签）
2. 包含内联CSS样式
3. 响应式设计
4. 直接返回HTML代码，不要任何解释""",
        
        "python": f"""请生成一个{req.theme}主题的Python代码。
要求：
1. 完整可运行的Python脚本
2. 良好的代码结构和中文注释
3. 直接返回代码，不要解释""",
        
        "word": f"""请生成一篇关于「{req.theme}」主题的Word文档内容。
要求：
1. 字数约{length_desc}
2. 包含标题、多个章节、结论
3. 结构清晰，段落分明
4. 直接返回文档内容，格式如下：
【标题】
章节一标题
章节一内容...

结论
结论内容...
不要使用Markdown格式，直接用纯文本""",
        
        "ppt": f"""请为「{req.theme}」主题生成一个PPT演示文稿的大纲。
要求：
1. 生成8-15页幻灯片内容
2. 格式如下：
第1页：封面 - {req.theme}
第2页：目录
...
直接返回大纲内容，每页用一行表示""",
        
        "document": f"""请生成一份关于「{req.theme}」的{req.task}文档。
要求：
1. 字数约{length_desc}
2. 专业文档格式
3. 包含摘要、正文、结论
4. 直接返回文档内容，不要格式标记""",
        
        "code": f"""请生成{req.task}的代码。
要求：
1. 完整可运行
2. 中文注释
3. 直接返回代码，不要解释""",
    }
    
    prompt = prompts.get(req.content_type, prompts["document"])
    
    _content_keywords = ['页面', '网站', '论文', '文章', '代码', '程序', '设计', '报告', '方案', '系统', '项目', '关于', '概述', '介绍', '总结', '分析', '研究']
    is_simple_text = (
        len(req.task) <= 10
        and not any(kw in req.task for kw in _content_keywords)
    )
    
    if is_simple_text and req.content_type in ("document", "word"):
        content = req.task
        print(f"[content] 简单文本直接写入: task='{req.task}', len={len(req.task)}")
    else:
        print(f"[content] 调用LLM生成: task='{req.task}', content_type={req.content_type}, is_simple={is_simple_text}")
        content = await call_llm(prompt)
    
    # 检查LLM是否失败
    if content.startswith("LLM"):
        return GenerateResponse(
            success=False,
            content=content,
            content_type=req.content_type,
            message="内容生成失败",
            error_detail=content
        )
    
    # 解析输出路径
    output_dir, explicit_file = _resolve_output_path(req)
    os.makedirs(output_dir, exist_ok=True)
    
    # 确定最终文件路径
    filename = _get_default_filename(req.content_type, explicit_file)
    final_path = explicit_file if explicit_file else os.path.join(output_dir, filename)
    
    # 保存文件
    output_file = None
    save_error = None
    
    if req.content_type in ["word", "document"]:
        output_file, save_error = save_word_document(final_path, content, title=req.theme)
    else:
        output_file, save_error = save_file(final_path, content)
    
    # 自动打开
    file_opened = False
    browser_opened = False
    if output_file and req.auto_open:
        if req.content_type == "html":
            browser_url = f"file:///{output_file.replace(os.sep, '/')}"
            browser_opened = open_in_browser(browser_url)
            file_opened = browser_opened
        else:
            file_opened = open_file(output_file)
    
    # 浏览器（非HTML的额外浏览器打开需求）
    if req.browser_url and req.browser_url != "auto_open":
        browser_opened = open_in_browser(req.browser_url)
    
    # 构建消息
    if output_file:
        msg = f"生成成功，已保存至: {output_file}"
        if file_opened:
            msg += "，已自动打开"
    else:
        msg = f"内容生成成功但保存失败: {save_error or '未知错误'}"
    
    return GenerateResponse(
        success=True,
        content=content,
        content_type=req.content_type,
        output_file=output_file,
        browser_opened=browser_opened,
        message=msg,
        error_detail=save_error
    )


@router.post("/browser/generate")
async def browser_generate_content(req: GenerateRequest):
    """打开浏览器，处理网页内容后生成结果"""
    
    if not req.browser_url:
        raise HTTPException(status_code=400, detail="需要提供browser_url参数")
    
    # 1. 打开浏览器
    browser_opened = open_in_browser(req.browser_url)
    
    # 2. 如果有生成任务，也执行生成
    if req.task:
        result = await generate_content(req)
        return {
            "success": result.success,
            "browser_opened": browser_opened,
            "url": req.browser_url,
            "content": result.content,
            "output_file": result.output_file,
            "message": result.message,
            "error_detail": result.error_detail
        }
    
    return {
        "success": True,
        "browser_opened": browser_opened,
        "url": req.browser_url,
        "message": f"已在浏览器中打开: {req.browser_url}"
    }